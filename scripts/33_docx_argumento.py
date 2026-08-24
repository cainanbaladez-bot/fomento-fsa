# -*- coding: utf-8 -*-
"""
33_docx_argumento.py — DOCX do ARGUMENTO CONDENSADO (só a linha de raciocínio).

Pedido do Cainan (24/08/2026): um Word que seja apenas o argumento, sem os dados e
sem a metodologia, começando pela conclusão — o fomento precisa usar evidência para
consolidar as estratégias das chamadas e usar o seu poder indutor para reorganizar
produção e distribuição. Máximo de três páginas.

REGRA DESTE DOCUMENTO (correção do Cainan, 24/08): o texto é **dele**, não meu. A
fonte `site_src/argumento.md` é montada com frases LITERAIS do `site_src/ensaio.md`,
na ordem que arma o raciocínio. A única edição permitida é pontual, para alinhavar:
cortar a oração que traz número ou método, tirar o conector que ficou órfão depois do
corte, desfazer parêntese. Não reescrever, não resumir com palavras minhas.
A primeira versão foi reescrita na minha voz e recusada — não repetir o erro.

Estilo do documento: o mesmo do `scripts/32` (Calibri, títulos em azul-marinho), sem
figura nenhuma — é texto corrido de ponta a ponta.

Saída: textos-docx/<data>_ARGUMENTO.docx
Rodar:  .\\.venv\\Scripts\\python.exe scripts\\33_docx_argumento.py
"""
import os
import re
import sys
import datetime as dt

sys.stdout.reconfigure(encoding='utf-8')
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'site_src', 'argumento.md')
OUT = os.path.join(BASE, 'textos-docx', f'{dt.date.today().isoformat()}_ARGUMENTO.docx')

from docx import Document                       # noqa: E402
from docx.shared import Pt, Inches, RGBColor    # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH   # noqa: E402

NAVY, GREY, INK = RGBColor(0x2B, 0x35, 0x66), RGBColor(0x6B, 0x72, 0x80), RGBColor(0x1A, 0x1A, 0x1A)

# ── fonte ─────────────────────────────────────────────────────────────────────
meta, secs, cur = {'title': '', 'subtitle': '', 'eyebrow': ''}, [], None
for ln in open(SRC, encoding='utf-8').read().splitlines():
    if ln.startswith('@title '):
        meta['title'] = ln[7:].strip()
    elif ln.startswith('@subtitle '):
        meta['subtitle'] = ln[10:].strip()
    elif ln.startswith('@eyebrow '):
        meta['eyebrow'] = ln[9:].strip()
    elif ln.startswith('## '):
        cur = {'titulo': ln[3:].strip(), 'paras': []}
        secs.append(cur)
    elif ln.strip() and cur is not None:
        cur['paras'].append(ln.strip())

strip_md = lambda t: t.replace('**', '').replace('*', '').replace('`', '')
INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+\*)')

doc = Document()
for _sec in doc.sections:
    _sec.top_margin = _sec.bottom_margin = Inches(0.8)
    _sec.left_margin = _sec.right_margin = Inches(0.95)
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(7)
st.paragraph_format.line_spacing = 1.15


def add_runs(p, text):
    for tok in INLINE.split(text.replace('`', '')):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            p.add_run(tok[2:-2]).bold = True
        elif tok.startswith('*') and tok.endswith('*'):
            p.add_run(tok[1:-1]).italic = True
        else:
            p.add_run(tok)


# ── capa ──
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run(meta['eyebrow'].upper())
r.bold = True
r.font.size = Pt(9.5)
r.font.color.rgb = NAVY

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run(strip_md(meta['title']))
r.bold = True
r.font.size = Pt(21)
r.font.color.rgb = INK

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
r = p.add_run(strip_md(meta['subtitle']))
r.italic = True
r.font.size = Pt(12)
r.font.color.rgb = GREY

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
p.paragraph_format.left_indent = Inches(0.02)
r = p.add_run(
    'Este documento traz apenas o encadeamento do raciocínio. Os números, as fontes, os '
    'recortes e os limites de cada conta estão na versão completa, publicada em '
    'cainanbaladez-bot.github.io/fomento-fsa, onde cada afirmação abre o gráfico que a '
    'sustenta. Trabalho em andamento.')
r.font.size = Pt(10)
r.italic = True
r.font.color.rgb = GREY

# ── corpo ──
for i, s in enumerate(secs):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4 if i == 0 else 15)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(strip_md(s['titulo']))
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY
    for tx in s['paras']:
        pp = doc.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(pp, tx)

doc.save(OUT)
n_par = sum(len(s['paras']) for s in secs)
palavras = sum(len(t.split()) for s in secs for t in s['paras'])
print(f'OK → {os.path.relpath(OUT, BASE)} ({os.path.getsize(OUT) / 1024:.0f} KB)')
print(f'   {len(secs)} seções · {n_par} parágrafos · {palavras} palavras')
