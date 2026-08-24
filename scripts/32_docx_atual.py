# -*- coding: utf-8 -*-
"""
32_docx_atual.py — DOCX idêntico ao site, com o gráfico de cada trecho embutido.

Diferença para o `scripts/30_docx_gerar.py` (o Word de trabalho, com as figuras
antigas do painel e uma por pergunta): aqui as figuras são **as 34 visualizações de
trecho** (scripts/23 e 24), cada uma inserida logo depois do parágrafo que a ancora
no site — a mesma ligação que no HTML abre no hover. O texto sai do
`site_src/ensaio.md` sem alteração nenhuma, então o documento é o gêmeo impresso de
https://cainanbaladez-bot.github.io/fomento-fsa/ensaio.html

As âncoras vêm dos dois formatos que o projeto usa:
  · `hoverfigs.json`  (Parte I) — âncora textual: acha o parágrafo que contém `inicio`
  · `hoverfigs2.json` (Partes II–IV) — âncora por parágrafo: `secao` + `par`, a mesma
    numeração de blocos do `scripts/21`

Os PNG são renderizados por kaleido a partir dos specs Plotly e ficam em
`outputs/figs_docx/hover/` (reaproveitados entre execuções).

Saída: textos-docx/<data>_ATUAL-com-graficos.docx
Rodar:  .\\.venv\\Scripts\\python.exe scripts\\32_docx_atual.py
"""
import os
import re
import sys
import json
import datetime as dt

import plotly.io as pio

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import site_base as S  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'site_src', 'ensaio.md')
BASES = os.path.join(BASE, 'outputs', 'bases')
FIGDIR = os.path.join(BASE, 'outputs', 'figs_docx', 'hover')
NUVEM_JPG = os.path.join(BASE, 'textos-docx', 'nuvem-dados_o-registro-publico.jpg')
OUT = os.path.join(BASE, 'textos-docx',
                   f'{dt.date.today().isoformat()}_ATUAL-com-graficos.docx')
os.makedirs(FIGDIR, exist_ok=True)

# ── as 34 visualizações de trecho ─────────────────────────────────────────────
HOV = {}
for nome in ('hoverfigs.json', 'hoverfigs2.json'):
    p = os.path.join(BASES, nome)
    if os.path.exists(p):
        HOV.update(json.load(open(p, encoding='utf-8')))
if not HOV:
    sys.exit('! hoverfigs*.json ausentes — rode scripts/23 e 24 antes deste')


def png_de(gid, v):
    """Renderiza o spec para PNG (cache por arquivo já existente)."""
    caminho = os.path.join(FIGDIR, f'{gid}.png')
    if os.path.exists(caminho):
        return caminho
    fig = pio.from_json(json.dumps(v['spec']))
    alt = (fig.layout.height or 400)
    fig.update_layout(width=1180, height=int(alt * 1.12), margin=dict(t=54, b=54))
    fig.write_image(caminho, scale=2)
    return caminho


print(f'renderizando {len(HOV)} gráficos (kaleido)…')
PNG = {}
for i, (gid, v) in enumerate(HOV.items(), 1):
    PNG[gid] = png_de(gid, v)
    print(f'  {i:>2}/{len(HOV)}  {gid}')

# ── parse do md, com a MESMA numeração de blocos do scripts/21 ────────────────
raw = open(SRC, encoding='utf-8').read()
meta, epis, body = {'title': '', 'subtitle': '', 'eyebrow': ''}, [], []
for ln in raw.splitlines():
    if ln.startswith('@title '):
        meta['title'] = ln[7:].strip()
    elif ln.startswith('@subtitle '):
        meta['subtitle'] = ln[10:].strip()
    elif ln.startswith('@eyebrow '):
        meta['eyebrow'] = ln[9:].strip()
    elif ln.startswith('@epi '):
        q, _, a = ln[5:].rpartition('|')
        epis.append((q.strip(), a.strip()))
    else:
        body.append(ln)

secs, cur = [], None
for ln in body:
    m = re.match(r'##\s*\[(\w+)\]\s*(.+)$', ln)
    if m:
        cur = {'sid': m.group(1), 'title': m.group(2).strip(), 'raw': []}
        secs.append(cur)
    elif cur is not None:
        cur['raw'].append(ln)

# (secao, indice_do_bloco) → [gid, …]
ANCORA = {}
for gid, v in HOV.items():
    if v.get('par') is not None:
        ANCORA.setdefault((v['secao'], int(v['par'])), []).append(gid)
    else:                                    # âncora textual (Parte I)
        alvo = v.get('inicio', '')
        achou = False
        for s in secs:
            idx = -1
            for ln in s['raw']:
                if ln.strip():
                    idx += 1
                if alvo and alvo in ln:
                    ANCORA.setdefault((s['sid'], idx), []).append(gid)
                    achou = True
                    break
            if achou:
                break
        if not achou:
            print(f'  ! âncora textual não encontrada: {gid}')

# ── blocos, preservando o índice de cada um ──────────────────────────────────
def blocos(lines):
    out, para, lista, pini = [], [], [], [None]
    idx = -1

    def flush_p():
        if para:
            out.append(('p', ' '.join(para), pini[0]))
            para.clear()
            pini[0] = None

    def flush_l():
        for it, i in lista:
            out.append(('li', it, i))
        lista.clear()

    for ln in lines:
        s = ln.rstrip()
        if s.strip():
            idx += 1
        if s.strip() == '[NUVEM]':
            flush_p(); flush_l(); out.append(('nuvem', '', idx))
        elif not s.strip():
            flush_p(); flush_l()
        elif s.lstrip().startswith('### '):
            flush_p(); flush_l(); out.append(('sub', s.lstrip()[4:], idx))
        elif s.lstrip().startswith('- '):
            flush_p(); lista.append((s.lstrip()[2:], idx))
        else:
            flush_l()
            if pini[0] is None:
                pini[0] = idx
            para.append(s.strip())
    flush_p(); flush_l()
    return out


for s in secs:
    s['blocks'] = blocos(s['raw'])

strip_md = lambda t: t.replace('**', '').replace('*', '').replace('`', '')

# ══════════════════════════ DOCX ══════════════════════════
from docx import Document                       # noqa: E402
from docx.shared import Pt, Inches, RGBColor    # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH   # noqa: E402

NAVY, GREY, INK = RGBColor(0x2B, 0x35, 0x66), RGBColor(0x6B, 0x72, 0x80), RGBColor(0x1A, 0x1A, 0x1A)
IMG_W = Inches(6.1)

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(8)
st.paragraph_format.line_spacing = 1.18

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+\*)')


def add_runs(p, text):
    text = text.replace('`', '')
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            p.add_run(tok[2:-2]).bold = True
        elif tok.startswith('*') and tok.endswith('*'):
            p.add_run(tok[1:-1]).italic = True
        else:
            p.add_run(tok)


def body_p(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, text)
    return p


def heading(text, size, color=NAVY, before=14, after=6, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def figura(gid):
    """A imagem do trecho + título, legenda e fonte, como no popup do site."""
    v = HOV[gid]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(PNG[gid], width=IMG_W)

    pc = doc.add_paragraph()
    pc.paragraph_format.left_indent = Inches(0.25)
    pc.paragraph_format.right_indent = Inches(0.25)
    pc.paragraph_format.space_after = Pt(3)
    rt = pc.add_run('▸ ' + v['titulo'] + '  ')
    rt.bold = True
    rt.font.size = Pt(9.5)
    rt.font.color.rgb = NAVY
    rl = pc.add_run(v['legenda'])
    rl.font.size = Pt(9)
    rl.font.color.rgb = GREY

    pf = doc.add_paragraph()
    pf.paragraph_format.left_indent = Inches(0.25)
    pf.paragraph_format.space_after = Pt(14)
    rf = pf.add_run('Fonte: ' + v['fonte'])
    rf.italic = True
    rf.font.size = Pt(8)
    rf.font.color.rgb = GREY


# ── capa ──
pt = doc.add_paragraph()
pt.paragraph_format.space_after = Pt(2)
r = pt.add_run(meta['eyebrow'])
r.bold = True
r.font.size = Pt(10)
r.font.color.rgb = NAVY
pt2 = doc.add_paragraph()
pt2.paragraph_format.space_after = Pt(6)
r = pt2.add_run(strip_md(meta['title']))
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = INK
ps = doc.add_paragraph()
ps.paragraph_format.space_after = Pt(14)
r = ps.add_run(strip_md(meta['subtitle']))
r.italic = True
r.font.size = Pt(11.5)
r.font.color.rgb = GREY
nota = doc.add_paragraph()
nota.paragraph_format.space_after = Pt(16)
rn = nota.add_run(
    f'Versão em Word do texto publicado em cainanbaladez-bot.github.io/fomento-fsa/ensaio.html, '
    f'gerada em {dt.datetime.now():%d/%m/%Y}. O texto é idêntico ao do site, sem uma vírgula de '
    f'diferença. As {len(HOV)} visualizações que no site abrem ao passar o mouse na passagem '
    f'sublinhada aparecem aqui como imagem, logo depois do parágrafo a que pertencem, com a mesma '
    f'legenda e a mesma fonte. Trabalho em andamento: números e recortes ainda podem mudar.')
rn.font.size = Pt(10)
rn.font.color.rgb = GREY
rn.italic = True
for q, a in epis:
    pe = doc.add_paragraph()
    pe.paragraph_format.left_indent = Inches(0.35)
    pe.paragraph_format.space_after = Pt(6)
    rq = pe.add_run('“' + strip_md(q.strip('“”')) + '”')
    rq.italic = True
    rq.font.size = Pt(10.5)
    ra = pe.add_run('   — ' + a)
    ra.font.size = Pt(9.5)
    ra.font.color.rgb = GREY

# ── corpo ──
usadas, n_par = set(), 0
for s in secs:
    sid, title, blocks = s['sid'], s['title'], s['blocks']
    if re.match(r'parte\d$', sid):
        heading(strip_md(title), 17, page_break=True, before=0)
        for kind, tx, _i in blocks:
            body_p(tx)
        continue
    m = re.match(r'c(\d+)$', sid)
    if m:
        rot = S.CLAIM_BY_ID.get(sid, (None, None, ''))[2]
        heading(f'Pergunta {m.group(1)} — {rot}', 14, before=18, after=2)
        lead = doc.add_paragraph()
        lead.paragraph_format.space_after = Pt(6)
        rl = lead.add_run(strip_md(title))
        rl.bold = True
        rl.font.size = Pt(12.5)
        rl.font.color.rgb = NAVY
    elif sid != 'abertura':
        heading(strip_md(title), 16, page_break=True, before=14)

    for kind, tx, i in blocks:
        if kind == 'nuvem':
            if os.path.exists(NUVEM_JPG):
                pi = doc.add_paragraph()
                pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pi.paragraph_format.space_after = Pt(12)
                pi.add_run().add_picture(NUVEM_JPG, width=IMG_W)
            continue
        if kind == 'sub':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(strip_md(tx))
            r.bold = True
            r.underline = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = NAVY
            continue
        if kind == 'li':
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, tx)
        else:
            body_p(tx)
            n_par += 1
        for gid in ANCORA.get((sid, i), []):
            figura(gid)
            usadas.add(gid)

doc.save(OUT)
faltou = [g for g in HOV if g not in usadas]
print(f'\nOK → {os.path.relpath(OUT, BASE)} ({os.path.getsize(OUT) / 1024 / 1024:.1f} MB)')
print(f'   {n_par} parágrafos · {len(usadas)}/{len(HOV)} gráficos inseridos'
      + (f' | sem âncora: {", ".join(faltou)}' if faltou else ''))
